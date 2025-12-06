import os
from langchain_community.tools import ReadFileTool, WriteFileTool, ListDirectoryTool
from langchain_core.tools import tool
from typing import Optional

# PDF Libraries
try:
    import pdfplumber
    from pypdf import PdfReader
    PDF_SUPPORT = True
except ImportError:
    PDF_SUPPORT = False
    print("Warning: PDF support not available. Install 'pypdf' and 'pdfplumber'")

# Word Document Library
try:
    from docx import Document
    from docx.shared import Pt
    DOCX_SUPPORT = True
except ImportError:
    DOCX_SUPPORT = False
    print("Warning: Word document support not available. Install 'python-docx'")


WORKSPACE = "."

# ============== Text File Tools ==============
read_file_tool = ReadFileTool(root_dir=WORKSPACE)
write_file_tool = WriteFileTool(root_dir=WORKSPACE)
list_directory_tool = ListDirectoryTool(root_dir=WORKSPACE)


# ============== PDF Tools ==============
@tool
def read_pdf_tool(file_path: str) -> str:
    """
    Read text content from a PDF file.
    
    Args:
        file_path: Path to the PDF file (relative to workspace or absolute)
    
    Returns:
        Extracted text content from the PDF
    """
    if not PDF_SUPPORT:
        return "Error: PDF support not available. Install 'pypdf' and 'pdfplumber'"
    
    full_path = os.path.join(WORKSPACE, file_path) if not os.path.isabs(file_path) else file_path
    
    if not os.path.exists(full_path):
        return f"Error: File '{full_path}' not found"
    
    try:
        text_content = []
        with pdfplumber.open(full_path) as pdf:
            for page_num, page in enumerate(pdf.pages, 1):
                page_text = page.extract_text()
                if page_text:
                    text_content.append(f"--- Page {page_num} ---\n{page_text}")
        
        return "\n\n".join(text_content) if text_content else "No text content found in PDF"
    except Exception as e:
        return f"Error reading PDF: {str(e)}"


# ============== Word Document Tools ==============
@tool
def read_docx_tool(file_path: str) -> str:
    """
    Read text content from a Word document (.docx).
    
    Args:
        file_path: Path to the Word document (relative to workspace or absolute)
    
    Returns:
        Extracted text content from the document
    """
    if not DOCX_SUPPORT:
        return "Error: Word document support not available. Install 'python-docx'"
    
    full_path = os.path.join(WORKSPACE, file_path) if not os.path.isabs(file_path) else file_path
    
    if not os.path.exists(full_path):
        return f"Error: File '{full_path}' not found"
    
    try:
        doc = Document(full_path)
        paragraphs = []
        
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)
        
        # Also extract text from tables if any
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        
        return "\n\n".join(paragraphs) if paragraphs else "No text content found in document"
    except Exception as e:
        return f"Error reading Word document: {str(e)}"


@tool
def write_docx_tool(file_path: str, text: str) -> str:
    """
    Write text content to a Word document (.docx).
    
    Args:
        file_path: Path to the Word document (relative to workspace or absolute)
        text: Text content to write to the document
    
    Returns:
        Success or error message
    """
    if not DOCX_SUPPORT:
        return "Error: Word document support not available. Install 'python-docx'"
    
    full_path = os.path.join(WORKSPACE, file_path) if not os.path.isabs(file_path) else file_path
    
    try:
        doc = Document()
        
        # Split text by double newlines to create paragraphs
        paragraphs = text.split("\n\n")
        
        for para_text in paragraphs:
            if para_text.strip():
                # Check if it's a heading (starts with #)
                if para_text.startswith("# "):
                    doc.add_heading(para_text[2:].strip(), level=1)
                elif para_text.startswith("## "):
                    doc.add_heading(para_text[3:].strip(), level=2)
                elif para_text.startswith("### "):
                    doc.add_heading(para_text[4:].strip(), level=3)
                else:
                    # Handle code blocks
                    if para_text.startswith("```"):
                        lines = para_text.split("\n")
                        code_content = "\n".join(lines[1:-1] if lines[-1] == "```" else lines[1:])
                        p = doc.add_paragraph()
                        run = p.add_run(code_content)
                        run.font.name = "Consolas"
                        run.font.size = Pt(10)
                    else:
                        doc.add_paragraph(para_text.strip())
        
        doc.save(full_path)
        return f"Successfully wrote content to '{file_path}'"
    except Exception as e:
        return f"Error writing Word document: {str(e)}"


# ============== Universal File Handler ==============
@tool
def read_any_file(file_path: str) -> str:
    """
    Read content from any supported file type (txt, pdf, docx).
    Automatically detects file type and uses appropriate reader.
    
    Args:
        file_path: Path to the file
    
    Returns:
        File content as text
    """
    extension = os.path.splitext(file_path)[1].lower()
    
    if extension == ".pdf":
        return read_pdf_tool.invoke(file_path)
    elif extension == ".docx":
        return read_docx_tool.invoke(file_path)
    elif extension in [".txt", ".md", ".json", ".py", ".js", ".html", ".css"]:
        return read_file_tool.run(file_path)
    else:
        return f"Unsupported file type: {extension}. Supported: .txt, .pdf, .docx, .md, .json"


@tool
def write_any_file(file_path: str, text: str) -> str:
    """
    Write content to any supported file type (txt, docx).
    Automatically detects file type and uses appropriate writer.
    Note: PDF writing is not supported (read-only).
    
    Args:
        file_path: Path to the file
        text: Content to write
    
    Returns:
        Success or error message
    """
    extension = os.path.splitext(file_path)[1].lower()
    
    if extension == ".pdf":
        return "Error: PDF writing is not supported. PDFs are read-only. Consider using .docx instead."
    elif extension == ".docx":
        return write_docx_tool.invoke({"file_path": file_path, "text": text})
    elif extension in [".txt", ".md", ".json", ".py", ".js", ".html", ".css"]:
        return write_file_tool.run({"file_path": file_path, "text": text})
    else:
        return f"Unsupported file type: {extension}. Supported for writing: .txt, .docx, .md, .json"


# Export all tools
tools = {
    # Text file tools
    "read_file_tool": read_file_tool,
    "write_file_tool": write_file_tool,
    "list_directory_tool": list_directory_tool,
    # PDF tools
    "read_pdf_tool": read_pdf_tool,
    # Word document tools
    "read_docx_tool": read_docx_tool,
    "write_docx_tool": write_docx_tool,
    # Universal tools
    "read_any_file": read_any_file,
    "write_any_file": write_any_file,
}